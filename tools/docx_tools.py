"""Editable Word (.docx) Campaign Strategy & Creative Brief Generator (UC4)."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

import docx
from docx.shared import Inches, Pt, RGBColor

from config.brand_guidelines import BrandTheme, parse_brand_theme
from config.settings import get_settings


def generate_campaign_brief_docx(
    campaign_name: str,
    master_strapline: str = "TWO DISTINCT PATHWAYS. ONE BALANCED FORCE.",
    anchor_look_and_feel: str = "Look & Feel 1: Synchronized Tandem",
    executive_summary: str = (
        "Comprehensive omnichannel campaign strategy synthesizing grounded scientific evidence, "
        "custom brand look-and-feel directions, and multi-format creative deliverables."
    ),
    key_messages: Optional[List[str]] = None,
    embedded_image_paths: Optional[List[str]] = None,
    theme_prompt: str = "astrazeneca_light",
    output_filename: str = "campaign_creative_brief.docx",
) -> Dict[str, Any]:
    """Generate an editable Word (.docx) Campaign Strategy & Creative Brief with embedded visuals."""
    settings = get_settings()
    theme: BrandTheme = parse_brand_theme(theme_prompt)
    doc = docx.Document()

    p_r, p_g, p_b = theme.primary_rgb
    s_r, s_g, s_b = theme.secondary_rgb

    title = doc.add_heading(f"{campaign_name} — Campaign Strategy & Creative Brief", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(p_r, p_g, p_b)

    p_strap = doc.add_paragraph()
    r_label = p_strap.add_run("Single Master Strapline: ")
    r_label.bold = True
    r_val = p_strap.add_run(f'"{master_strapline}"')
    r_val.bold = True
    r_val.font.size = Pt(14)
    r_val.font.color.rgb = RGBColor(s_r, s_g, s_b)

    doc.add_heading("1. Executive Summary & Strategic Rationale", level=1)
    doc.add_paragraph(executive_summary)

    doc.add_heading("2. Recommended Anchor Look & Feel & Brand Theme", level=1)
    doc.add_paragraph(f"Recommended Anchor Direction: {anchor_look_and_feel}")
    doc.add_paragraph(
        f"Active Theme: {theme.theme_name} | Background: {theme.background_hex} | "
        f"Text: {theme.text_primary_hex} | Palette: {', '.join(theme.palette_hex)}"
    )

    doc.add_heading("3. Core Scientific & Commercial Messages", level=1)
    msgs = key_messages or [
        "Complementary dual-pathway synergy delivering broader clinical and patient value.",
        "Unified visual identity and single master strapline across slides, brochures, and video.",
        "Grounded scientific claims backed by live Google Search and uploaded source materials.",
    ]
    for m in msgs:
        doc.add_paragraph(m, style="List Bullet")

    if embedded_image_paths:
        doc.add_heading("4. Generated Campaign Visuals & Charts", level=1)
        for img_p in embedded_image_paths:
            if img_p and Path(img_p).exists():
                try:
                    doc.add_picture(str(img_p), width=Inches(5.8))
                except Exception:
                    pass

    out_path = settings.output_dir / output_filename
    doc.save(str(out_path))
    return {
        "status": "success",
        "docx_path": str(out_path),
        "campaign_name": campaign_name,
        "master_strapline": master_strapline,
        "theme_id": theme.theme_id,
    }
